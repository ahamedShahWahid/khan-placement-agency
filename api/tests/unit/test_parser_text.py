"""Unit tests for extract_text() — uses fpdf2 + python-docx to generate fixtures in-process."""

from __future__ import annotations

import io

import pytest
from docx import Document
from fpdf import FPDF
from pypdf import PdfReader, PdfWriter

from kpa.integrations.parser.base import ParserError
from kpa.integrations.parser.text import MAX_TEXT_BYTES, extract_text

# --- Fixture generators (in-process; no binary commits) ---


def _make_pdf(text_lines: list[str]) -> bytes:
    """Generate a minimal PDF with the given text lines."""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in text_lines:
        pdf.cell(text=line, new_x="LMARGIN", new_y="NEXT")
    out = pdf.output()
    return bytes(out)


def _make_docx(text_lines: list[str]) -> bytes:
    """Generate a minimal DOCX with the given paragraphs."""
    doc = Document()
    for line in text_lines:
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_password_protected_pdf(text: str, password: str) -> bytes:
    src = _make_pdf([text])
    reader = PdfReader(io.BytesIO(src))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(user_password=password)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _make_blank_pdf() -> bytes:
    """PDF with a page but no text content — simulates an image-only resume."""
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


# --- Tests ---


async def test_extract_text_from_pdf() -> None:
    # Lines are long enough (> _EMPTY_THRESHOLD=50 chars total) so pypdf returns them
    # directly without triggering the pdfminer fallback.
    line1 = "Hello world, this is a resume text line for extraction testing."
    line2 = "Second line contains more text to ensure threshold is met properly."
    pdf_bytes = _make_pdf([line1, line2])
    text = await extract_text(content=pdf_bytes, content_type="application/pdf")
    assert "Hello world" in text
    assert "Second line" in text


async def test_extract_text_from_docx() -> None:
    docx_bytes = _make_docx(["Hello world", "Second paragraph"])
    text = await extract_text(
        content=docx_bytes,
        content_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )
    assert "Hello world" in text
    assert "Second paragraph" in text


async def test_extract_text_rejects_legacy_doc() -> None:
    with pytest.raises(ParserError, match="doc_legacy_not_supported"):
        await extract_text(content=b"\xd0\xcf\x11\xe0", content_type="application/msword")


async def test_extract_text_rejects_unsupported_content_type() -> None:
    with pytest.raises(ParserError, match="unsupported_content_type"):
        await extract_text(content=b"random", content_type="image/png")


async def test_extract_text_rejects_password_protected_pdf() -> None:
    pdf_bytes = _make_password_protected_pdf("secret content", password="abc")
    with pytest.raises(ParserError, match="password_protected"):
        await extract_text(content=pdf_bytes, content_type="application/pdf")


async def test_extract_text_rejects_image_only_pdf() -> None:
    """Blank page → no extractable text → both pypdf + pdfminer return empty → ParserError."""
    pdf_bytes = _make_blank_pdf()
    with pytest.raises(ParserError, match="no_text_extracted"):
        await extract_text(content=pdf_bytes, content_type="application/pdf")


async def test_extract_text_rejects_all_whitespace_docx() -> None:
    """A DOCX where every paragraph is blank → no_text_extracted.

    DOCX equivalent of the image-only PDF case.
    """
    docx_bytes = _make_docx(["   ", "\t", ""])
    with pytest.raises(ParserError, match="no_text_extracted"):
        await extract_text(
            content=docx_bytes,
            content_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
        )


async def test_extract_text_truncates_to_max_bytes() -> None:
    long_line = "x" * 100  # 100 bytes per line
    pdf_bytes = _make_pdf([long_line] * 1000)  # ~100 KB of text — over the 64KB cap
    text = await extract_text(content=pdf_bytes, content_type="application/pdf")
    assert len(text.encode("utf-8")) <= MAX_TEXT_BYTES
