"""Unit tests for LibraryResumeParser — exercises regex + dictionary logic.

Tests feed in canned PDF/DOCX bytes (generated via fpdf2 / python-docx) and
assert the ParsedResume shape. Heuristics for name/experience/education are
intentionally loose; tests assert what the regex CAN find, not what's correct.
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from fpdf import FPDF

from jobify.integrations.parser.base import ParsedResume
from jobify.integrations.parser.library import LibraryResumeParser, _extract_skills
from jobify.integrations.parser.skills_dict import SKILLS

PDF_CT = "application/pdf"
DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _pdf(text_lines: list[str]) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in text_lines:
        pdf.cell(text=line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def _docx(text_lines: list[str]) -> bytes:
    doc = Document()
    for line in text_lines:
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def parser() -> LibraryResumeParser:
    return LibraryResumeParser()


async def test_parse_returns_parsed_resume_with_parser_name(
    parser: LibraryResumeParser,
) -> None:
    pr = await parser.parse(content=_pdf(["hello"]), content_type=PDF_CT)
    assert isinstance(pr, ParsedResume)
    assert pr.parser_name == "library.v1"
    assert pr.schema_version == 1
    assert "hello" in pr.raw_text


async def test_parse_extracts_email(parser: LibraryResumeParser) -> None:
    pr = await parser.parse(
        content=_pdf(["Ahamed Wahid", "Contact: ahamed.wahid@example.com"]),
        content_type=PDF_CT,
    )
    assert pr.email == "ahamed.wahid@example.com"


async def test_parse_extracts_indian_phone(parser: LibraryResumeParser) -> None:
    pr = await parser.parse(
        content=_pdf(["Name: A", "Phone: +91-98765-43210"]), content_type=PDF_CT
    )
    assert pr.phone is not None
    # Stripped of separators, the digits should match the input.
    digits = "".join(ch for ch in pr.phone if ch.isdigit() or ch == "+")
    assert "9876543210" in digits


async def test_parse_extracts_intl_phone(parser: LibraryResumeParser) -> None:
    pr = await parser.parse(
        content=_pdf(["Name: A", "Mobile: +1 415 555 0123"]), content_type=PDF_CT
    )
    assert pr.phone is not None


async def test_parse_extracts_skills_from_dictionary(parser: LibraryResumeParser) -> None:
    pr = await parser.parse(
        content=_pdf(
            [
                "Senior Engineer",
                "Skills: Python, FastAPI, Postgres, Docker, Kubernetes",
            ]
        ),
        content_type=PDF_CT,
    )
    assert set(pr.skills) >= {"python", "fastapi", "postgres", "docker", "kubernetes"}


async def test_parse_skills_are_deduped_and_sorted(parser: LibraryResumeParser) -> None:
    pr = await parser.parse(
        content=_pdf(["python python PYTHON fastapi FastAPI", "more python"]),
        content_type=PDF_CT,
    )
    assert pr.skills == sorted(set(pr.skills))
    assert pr.skills.count("python") == 1


async def test_parse_finds_experience_date_ranges(parser: LibraryResumeParser) -> None:
    pr = await parser.parse(
        content=_pdf(
            [
                "Experience",
                "Acme Corp: Jan 2020 to Dec 2022 - built things",
                "BetaWorks: 2018 - 2020 - Senior Eng",
            ]
        ),
        content_type=PDF_CT,
    )
    assert len(pr.experience) >= 1
    # The regex should pick up at least one start/end pair.
    has_range = any(e.start is not None and e.end is not None for e in pr.experience)
    assert has_range


async def test_parse_finds_education_degree(parser: LibraryResumeParser) -> None:
    pr = await parser.parse(
        content=_pdf(["Education", "B.Tech, IIT Bombay, 2018"]), content_type=PDF_CT
    )
    assert any(e.degree is not None for e in pr.education)
    assert any(e.end_year == 2018 for e in pr.education)


async def test_parse_empty_resume_returns_valid_parsed_resume(
    parser: LibraryResumeParser,
) -> None:
    """A resume with only whitespace still produces a valid ParsedResume
    (no exceptions; empty arrays for everything except raw_text)."""
    pr = await parser.parse(content=_pdf(["   "]), content_type=PDF_CT)
    assert pr.email is None
    assert pr.phone is None
    assert pr.skills == []
    assert pr.experience == []
    assert pr.education == []


@pytest.mark.parametrize(
    ("text", "phantom"),
    [
        ("Senior Engineer with 6 years", "gin"),
        ("Team member since 2019", "ember"),
        ("Studied at Symbiosis Institute", "ios"),
        ("Handled escalation workflows", "scala"),
        ("Built a hyperlocal delivery app", "perl"),
        ("Reduced margin of error", "gin"),
    ],
)
def test_skills_ignores_substrings_inside_longer_words(text: str, phantom: str) -> None:
    """The dictionary is matched on token boundaries, not raw containment.

    Every `phantom` here is a real SKILLS entry that naive `skill in text`
    finds inside an unrelated English word. These were the documented
    false-positive family in `000_README.md`; each one costs precision on
    the gated `skills` F1, which counts FPs.
    """
    assert phantom in SKILLS, f"{phantom!r} must be a real dict entry for this to mean anything"
    assert phantom not in _extract_skills(text)


def test_skills_prefers_the_longer_overlapping_entry() -> None:
    """`postgres` and `sql` are both entries and both live inside `postgresql`.

    Containment reported all three from one token; boundary matching reports
    only the token actually written.
    """
    found = _extract_skills("Databases: PostgreSQL")
    assert "postgresql" in found
    assert "postgres" not in found
    assert "sql" not in found


@pytest.mark.parametrize(
    "skill",
    ["c++", "c#", "node.js", "asp.net", "ci/cd", "scikit-learn", "spring boot"],
)
def test_skills_still_matches_entries_with_punctuation(skill: str) -> None:
    """Regression guard for the boundary rule itself.

    A naive ``\\b…\\b`` breaks on entries whose first/last character is not a
    word character (`c++`, `c#`), because `\\b` is defined between a word and
    a non-word char — so it demands a word char adjacent to the `+`/`#` and
    the entry stops matching. The rule must anchor per-side on whether the
    entry's own edge is alphanumeric.
    """
    assert skill in SKILLS
    assert skill in _extract_skills(f"Skills: {skill}, python")


def test_skills_suppresses_an_entry_nested_in_a_longer_match() -> None:
    """ "Spring Boot" must not also report the bare "spring" entry."""
    found = _extract_skills("Stack: Java, Spring Boot, MySQL")
    assert "spring boot" in found
    assert "spring" not in found


def test_skills_keeps_a_short_entry_that_also_stands_alone() -> None:
    """Suppression is per-occurrence, not per-name.

    This resume writes BOTH "Postgres" on its own and "PostgreSQL"; the short
    entry is genuinely present, so it must survive. A name-based rule ("drop
    any entry that is a substring of another match") would wrongly delete it.
    """
    found = _extract_skills("Datastores: Postgres pgvector. Also PostgreSQL.")
    assert {"postgres", "postgresql", "pgvector"}.issubset(found)


def test_skills_matches_at_string_edges() -> None:
    """Boundaries at start/end of input must not require a neighbour char."""
    assert "python" in _extract_skills("python")


async def test_parse_works_on_docx(parser: LibraryResumeParser) -> None:
    pr = await parser.parse(
        content=_docx(["Name", "Email: a@b.com", "Skills: Java, Spring Boot"]),
        content_type=DOCX_CT,
    )
    assert pr.email == "a@b.com"
    assert {"java", "spring boot"}.issubset(set(pr.skills))


async def test_parse_finds_certification(parser: LibraryResumeParser) -> None:
    pr = await parser.parse(
        content=_pdf(["Certifications", "AWS Certified Solutions Architect, 2022"]),
        content_type=PDF_CT,
    )
    assert any(c.year == 2022 for c in pr.certifications)


async def test_parse_name_skips_section_headers(parser: LibraryResumeParser) -> None:
    """Section labels like 'Skills:' must not be mistaken for the candidate's name."""
    pr = await parser.parse(
        content=_pdf(["Skills:", "Python FastAPI", "Email: x@y.com"]),
        content_type=PDF_CT,
    )
    # No usable name in this fixture; certainly not "Skills:".
    assert pr.name != "Skills:"
